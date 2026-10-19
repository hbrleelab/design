"""Generate the HBRLRG conference poster template (PPTX).

    python3 tools/build_poster.py [outdir] [--format a0-portrait] [--fonts DIR]

A poster is read at two distances, and the type scale follows from them: the
title has to carry across an aisle (about 5 m), the body has to be read by
someone standing in front of it (about 1.5 m). At A0 that puts the title near
100 pt and body text at 30 pt — the floor, not a target. Nothing on the poster
goes below 20 pt.

Layout is a fixed column grid rather than free placement. Each column is a
stack of blocks (heading, text, figure, ...); one figure per column is "flex"
and absorbs whatever height is left, so every column ends flush on the same
line however much text a format holds. Students edit the boxes; the grid is
what keeps four authors' posters looking like they came from one group.

Text heights are measured, not guessed: every block is word-wrapped against
Archivo's real advance widths, so a box is exactly as tall as its text and the
next block starts where the text ends. That needs the font files — the same
ones the release workflow fetches into .fonts/.

Formats are entries in FORMATS. Only A0 portrait exists until the layout
references in references/posters/ say which others the group actually uses.
"""
import argparse
from pathlib import Path

from fontTools.ttLib import TTFont
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Mm, Pt

ROOT = Path(__file__).resolve().parent.parent
ASSETS = ROOT / "assets"

# --- design tokens (design-system/HBRLRG-design-system.md §1) --------------
NAVY = RGBColor(0x14, 0x24, 0x3F)
TEAL = RGBColor(0x0E, 0xA7, 0x9A)
INK = RGBColor(0x2B, 0x2F, 0x36)
MUTED = RGBColor(0x6B, 0x72, 0x80)
SURFACE = RGBColor(0xF6, 0xF7, 0xF9)
BORDER = RGBColor(0xE4, 0xE7, 0xEC)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
ON_NAVY = RGBColor(0x93, 0xA2, 0xBC)
LATIN, HANGUL = "Archivo", "Pretendard"

PT_MM = 25.4 / 72
# Office sets single spacing at 1.2 em for Archivo, and puts the extra of a
# proportional spacing above each line — the first one included. Measured on
# the LibreOffice export, which is what the sample PDF is.
SINGLE = 1.2
X_HEIGHT = 0.526

# --- formats ---------------------------------------------------------------
FORMATS = {
    "a0-portrait": dict(
        name="A0 세로", w=841, h=1189, margin=32, gutter=24, cols=3,
        footer_h=78,
        type=dict(title=100, authors=40, affil=24, eyebrow=24, heading=56,
                  body=30, caption=24, small=20),
    ),
}


# --- text measurement ------------------------------------------------------
class Measure:
    """Greedy word wrap on Archivo's advance widths. Hangul falls back to a
    flat 0.95 em (Pretendard's syllables are close to square); the wrap width
    keeps 3 % in reserve for kerning and for PowerPoint wrapping a hair
    earlier than LibreOffice."""

    def __init__(self, font_dir):
        self.faces = {}
        for bold, name in ((False, "Archivo-Regular.ttf"), (True, "Archivo-Bold.ttf")):
            f = TTFont(Path(font_dir) / name)
            self.faces[bold] = (f.getBestCmap(), f["hmtx"], f["head"].unitsPerEm)

    def width(self, s, size, bold=False):
        cmap, hmtx, upm = self.faces[bold]
        em = sum(hmtx[cmap[ord(c)]][0] / upm if ord(c) in cmap else 0.95 for c in s)
        return em * size * PT_MM

    def lines(self, s, w, size, bold=False):
        limit = w * 0.97
        n, cur = 1, ""
        for word in s.split(" "):
            trial = f"{cur} {word}" if cur else word
            if cur and self.width(trial, size, bold) > limit:
                n, cur = n + 1, word
            else:
                cur = trial
        return n


def find_fonts(explicit):
    """--fonts, then .fonts/ (where the release workflow puts them), then the
    usual per-user and system font folders."""
    candidates = [Path(explicit)] if explicit else [
        ROOT / ".fonts", Path.home() / ".local/share/fonts",
        Path.home() / "Library/Fonts", Path("/usr/share/fonts"),
    ]
    for c in candidates:
        hits = sorted(c.rglob("Archivo-Regular.ttf")) if c.is_dir() else []
        if hits and (hits[0].parent / "Archivo-Bold.ttf").exists():
            return hits[0].parent
    raise SystemExit(
        "Archivo-Regular.ttf / Archivo-Bold.ttf 를 찾지 못했습니다. --fonts 로 폴더를 "
        "지정하거나, release.yml 의 Fetch fonts 단계처럼 .fonts/ 에 받아 두세요.")


# --- low-level helpers -----------------------------------------------------
def line_pitch(size, spacing):
    return size * SINGLE * spacing * PT_MM


def first_baseline(size, spacing):
    """Distance from a text box's top to its first baseline, in mm."""
    return (SINGLE * spacing - 0.2) * size * PT_MM


def rect(slide, x, y, w, h, fill, line=None):
    s = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Mm(x), Mm(y), Mm(w), Mm(h))
    s.fill.solid()
    s.fill.fore_color.rgb = fill
    if line is None:
        s.line.fill.background()      # never read s.line.color: it adds <a:ln>
    else:
        s.line.fill.solid()
        s.line.fill.fore_color.rgb = line
        s.line.width = Pt(1.5)
    s.shadow.inherit = False
    # The theme style reference would still give LibreOffice a drop shadow
    # (it ignores the empty effectLst); every property it supplies is set
    # explicitly above, so drop it.
    style = s._element.find(qn("p:style"))
    if style is not None:
        s._element.remove(style)
    return s


def _run_font(run, size, colour, bold=False, tracking=None, sup=False):
    run.font.name = LATIN
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = colour
    rPr = run._r.get_or_add_rPr()
    rPr.append(rPr.makeelement(qn("a:ea"), {"typeface": HANGUL}))
    if tracking:
        rPr.set("spc", str(int(tracking * size * 100)))
    if sup:
        rPr.set("baseline", "30000")


def text(slide, x, y, w, h, runs, size, colour, *, bold=False, align=PP_ALIGN.LEFT,
         spacing=1.25, tracking=None, anchor=MSO_ANCHOR.TOP, name=None):
    """runs: str, or a list of paragraphs, each a str or a list of
    (text, {overrides}) pieces."""
    box = slide.shapes.add_textbox(Mm(x), Mm(y), Mm(w), Mm(h))
    if name:
        box.name = name
    tf = box.text_frame
    tf.word_wrap = True
    tf.auto_size = None
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    paras = [runs] if isinstance(runs, str) else runs
    for i, para in enumerate(paras):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.line_spacing = spacing
        pieces = [(para, {})] if isinstance(para, str) else para
        for t, o in pieces:
            r = p.add_run()
            r.text = t
            _run_font(r, o.get("size", size), o.get("colour", colour),
                      o.get("bold", bold), o.get("tracking", tracking), o.get("sup", False))
    return box


def picture(slide, path, x, y, h=None, w=None):
    """Set one dimension; the other follows the file."""
    from PIL import Image
    with Image.open(path) as im:
        ratio = im.width / im.height
    if h is not None:
        w = h * ratio
    else:
        h = w / ratio
    slide.shapes.add_picture(str(path), Mm(x), Mm(y), Mm(w), Mm(h))
    return w, h


# --- the poster ------------------------------------------------------------
class Poster:
    BODY_SPACING = 1.35
    GAP = 22                 # between sections
    AFTER_TEXT = 10          # text -> next block in the same section
    AFTER_FIGURE = 12        # caption -> next block
    BULLET_GAP = 3

    def __init__(self, fmt, measure):
        self.F = F = FORMATS[fmt]
        self.T = F["type"]
        self.m = measure
        W, M, G, C = F["w"], F["margin"], F["gutter"], F["cols"]
        self.col_w = (W - 2 * M - (C - 1) * G) / C
        self.col_x = [M + i * (self.col_w + G) for i in range(C)]

    # Every block is a pair: height(block) for the layout pass, draw(...) for
    # the drawing pass. Both read the same numbers, so they cannot disagree.
    def text_h(self, s, w, size, spacing):
        return self.m.lines(s, w, size) * line_pitch(size, spacing)

    def heading_h(self):
        return self.T["heading"] * PT_MM * 1.15 + 5 + 2.4 + 10

    def heading(self, s, x, y, num, title):
        """Number in teal, title in navy, then the teal-capped rule underneath —
        the same device as the documents and slides, one per section."""
        size = self.T["heading"]
        h = size * PT_MM * 1.15
        pieces = [(f"{num}  ", {"colour": TEAL})] if num else []
        text(s, x, y, self.col_w, h, [pieces + [(title, {})]], size, NAVY, bold=True,
             spacing=1.0, name=f"섹션 {num} 제목")
        ry = y + h + 5
        rect(s, x, ry + 0.8, self.col_w, 0.8, NAVY)
        rect(s, x, ry, 30, 2.4, TEAL)
        return self.heading_h()

    def body(self, s, x, y, w, str_, colour=INK, name="본문"):
        size = self.T["body"]
        h = self.text_h(str_, w, size, self.BODY_SPACING)
        text(s, x, y, w, h, str_, size, colour, spacing=self.BODY_SPACING, name=name)
        return h

    def bullets_h(self, items, w):
        size = self.T["body"]
        indent = self._indent(size)
        return sum(self.text_h(i, w - indent, size, self.BODY_SPACING) for i in items) \
            + self.BULLET_GAP * (len(items) - 1)

    def _indent(self, size):
        return size * PT_MM * 1.05

    def bullets(self, s, x, y, w, items, colour=INK, name="요점"):
        """Teal squares centred on the first line's x-height, as in the documents."""
        size = self.T["body"]
        sq = size * 0.3 * PT_MM
        indent = self._indent(size)
        mid = first_baseline(size, self.BODY_SPACING) - X_HEIGHT / 2 * size * PT_MM
        for i, item in enumerate(items):
            h = self.text_h(item, w - indent, size, self.BODY_SPACING)
            rect(s, x, y + mid - sq / 2, sq, sq, TEAL)
            text(s, x + indent, y, w - indent, h, item, size, colour,
                 spacing=self.BODY_SPACING, name=f"{name} {i + 1}")
            y += h + self.BULLET_GAP

    def caption_h(self, caption):
        size = self.T["caption"]
        return 6 + self.text_h(caption, self.col_w, size, 1.3)

    def figure(self, s, x, y, h, n, caption, aspect_label):
        """A figure well — where a plot or micrograph gets pasted in — with its
        caption. The label inside names the slot's shape and size so the
        author can prepare the image at the right proportions."""
        T = self.T
        box = rect(s, x, y, self.col_w, h, SURFACE, line=BORDER)
        box.name = f"그림 {n} 자리"
        tf = box.text_frame
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        p = tf.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER
        r = p.add_run()
        r.text = f"Figure {n}  ·  {aspect_label}  ·  {self.col_w:.0f} × {h:.0f} mm"
        _run_font(r, T["caption"], MUTED)
        ch = self.caption_h(caption) - 6
        text(s, x, y + h + 6, self.col_w, ch, caption, T["caption"], MUTED, spacing=1.3,
             name=f"그림 {n} 설명")

    # -- layout pass --------------------------------------------------------
    def height(self, b):
        kind, w = b[0], self.col_w
        if kind == "heading":
            return self.heading_h()
        if kind == "text":
            return self.text_h(b[1], w, self.T["body"], self.BODY_SPACING)
        if kind == "bullets":
            return self.bullets_h(b[1], w)
        if kind == "figure":
            return w / b[2] + self.caption_h(b[4])
        if kind == "flex":
            return self.caption_h(b[3])
        if kind == "conclusion":
            return self.heading_h() + 2 * 16 + self.bullets_h(b[3], w - 32)
        if kind == "references":
            size = self.T["small"]
            return self.label_h() + sum(self.text_h(r, w, size, 1.3) for r in b[1])
        raise ValueError(kind)

    def spaced(self, blocks):
        """(block, space after) pairs. A ("gap",) marker starts a new section
        and sets the space before it; the last block ends flush."""
        out = []
        for i, b in enumerate(blocks):
            if b[0] == "gap":
                continue
            nxt = blocks[i + 1] if i + 1 < len(blocks) else None
            if nxt is None:
                after = 0
            elif nxt[0] == "gap":
                after = self.GAP
            elif b[0] == "heading":
                after = 0                      # heading_h carries its own space
            elif b[0] in ("figure", "flex"):
                after = self.AFTER_FIGURE
            else:
                after = self.AFTER_TEXT
            out.append((b, after))
        return out

    def label_h(self):
        return self.T["small"] * PT_MM * 1.3 + 4

    def label(self, s, x, y, w, t):
        text(s, x, y, w, self.label_h(), t, self.T["small"], MUTED, bold=True,
             tracking=0.14, spacing=1.0, name=t)

    # -- drawing pass -------------------------------------------------------
    def column(self, s, x, top, bottom, blocks):
        blocks = self.spaced(blocks)
        n_flex = sum(1 for b, _ in blocks if b[0] == "flex")
        fixed = sum(self.height(b) + after for b, after in blocks)
        flex_h = (bottom - top - fixed) / max(n_flex, 1)
        if n_flex and flex_h < 100:
            raise SystemExit(f"그림 자리 높이 {flex_h:.0f} mm — 이 단에 글이 너무 많습니다")
        y = top
        for b, after in blocks:
            kind, h = b[0], self.height(b)
            if kind == "heading":
                self.heading(s, x, y, b[1], b[2])
            elif kind == "text":
                self.body(s, x, y, self.col_w, b[1])
            elif kind == "bullets":
                self.bullets(s, x, y, self.col_w, b[1])
            elif kind == "figure":
                self.figure(s, x, y, self.col_w / b[2], b[1], b[4], b[3])
            elif kind == "flex":
                self.figure(s, x, y, flex_h, b[1], b[3], b[2])
                h += flex_h
            elif kind == "conclusion":
                self.heading(s, x, y, b[1], b[2])
                hd = self.heading_h()
                rect(s, x, y + hd, self.col_w, h - hd, SURFACE).name = "결론 상자"
                self.bullets(s, x + 16, y + hd + 16, self.col_w - 32, b[3], NAVY, "결론")
            elif kind == "references":
                self.label(s, x, y, self.col_w, "REFERENCES")
                size = self.T["small"]
                text(s, x, y + self.label_h(), self.col_w, h - self.label_h(), b[1],
                     size, MUTED, spacing=1.3, name="참고문헌")
            y += h + after
        return y, flex_h


def build(fmt, out, font_dir):
    P = Poster(fmt, Measure(font_dir))
    F, T = P.F, P.T
    W, H, M = F["w"], F["h"], F["margin"]

    prs = Presentation()
    prs.slide_width, prs.slide_height = Mm(W), Mm(H)
    s = prs.slides.add_slide(prs.slide_layouts[6])

    # header ---------------------------------------------------------------
    # Built top-down and measured, so a three-line title pushes the authors
    # and the affiliation down instead of running into them.
    title = "Poster title in one or two lines, stated as the finding rather than the topic"
    title_h = P.m.lines(title, W - 2 * M, T["title"], bold=True) * line_pitch(T["title"], 0.98)
    authors = [("First Author", {}), ("1", {"sup": True}), (",  Second Author", {}),
               ("1", {"sup": True}), (",  Han-Bo-Ram Lee", {}), ("1,*", {"sup": True})]
    authors_h = line_pitch(T["authors"], 1.2)
    emblem = 30          # 3.5 × the affiliation size, as on the slide cover
    y_title = 82
    y_auth = y_title + title_h + 6
    y_affil = y_auth + authors_h + 10
    HH = y_affil + emblem + 24
    rect(s, 0, 0, W, HH, NAVY).name = "머리 배경"
    picture(s, ASSETS / "logo/hbrlrg-horizontal-reverse.png", M, 30, h=22)
    text(s, W - M - 420, 30, 420, 22, "CONFERENCE NAME 2026  ·  POSTER P-00",
         T["eyebrow"], TEAL, bold=True, align=PP_ALIGN.RIGHT, tracking=0.14,
         anchor=MSO_ANCHOR.MIDDLE, name="학회명 · 포스터 번호")
    rect(s, M, 70, 40, 3, TEAL).name = "청록 막대"
    text(s, M, y_title, W - 2 * M, title_h, title, T["title"], WHITE, bold=True,
         spacing=0.98, name="제목")
    text(s, M, y_auth, W - 2 * M, authors_h, [authors], T["authors"], WHITE, spacing=1.2,
         name="저자")
    # The emblem always travels with the affiliation it stands for, and the
    # text is centred on it (design system §7).
    ew, _ = picture(s, ASSETS / "unist-emblem-onnavy.png", M, y_affil, h=emblem)
    ax = M + ew + T["affil"] * PT_MM
    text(s, ax, y_affil, W - M - ax, emblem,
         [[("1 ", {"sup": True}),
           ("Graduate School of Semiconductor Materials & Devices Engineering, "
            "Ulsan National Institute of Science and Technology (UNIST)", {})],
          [("* Corresponding author  ·  hbrlee@unist.ac.kr", {})]],
         T["affil"], ON_NAVY, spacing=1.3, anchor=MSO_ANCHOR.MIDDLE, name="소속")

    # footer ---------------------------------------------------------------
    FH = F["footer_h"]
    fy = H - FH
    rect(s, 0, fy, W, FH, SURFACE).name = "꼬리 배경"
    qr = 52
    ty = fy + 14
    P.label(s, M, ty, 380, "CONTACT")
    text(s, M, ty + P.label_h(), 380, 40,
         ["First Author  ·  first.author@unist.ac.kr",
          [("HBRL Research Group  ·  https://hbrl-research.group",
            {"size": T["small"], "colour": MUTED})]],
         26, INK, spacing=1.35, name="연락처")
    ack_x = M + 400
    P.label(s, ack_x, ty, 300, "ACKNOWLEDGEMENTS")
    text(s, ack_x, ty + P.label_h(), W - M - qr - 20 - ack_x, 40,
         "This work was supported by the funding agency (grant no. 0000-0000). "
         "Replace with your own funding statement.",
         T["small"], MUTED, spacing=1.35, name="사사")
    picture(s, ASSETS / "qr/hbrl-research-group.png", W - M - qr, fy + (FH - qr) / 2, h=qr)

    # body -----------------------------------------------------------------
    top, bottom = HH + 32, fy - 32
    columns = [
        [("heading", "01", "Introduction"),
         ("text", "Two or three sentences: the problem, why it matters to the field, "
                  "and the gap this work fills. Body text stays at 30 pt — it is read "
                  "by someone standing in front of the poster, not leaning in."),
         ("bullets", ["State the question you answered",
                      "Give the one number that frames it",
                      "Say what was missing before"]),
         ("gap",),
         ("heading", "02", "Methods"),
         ("text", "Only what a visitor needs to trust the result: process, "
                  "precursors, conditions. The full recipe belongs in the paper."),
         ("flex", 1, "free", "Figure 1. Process schematic or reactor layout.")],
        [("heading", "03", "Results"),
         ("text", "Lead with the finding. One claim per figure, and the claim in "
                  "the caption, so the poster still reads with the text skipped."),
         ("figure", 2, 16 / 9, "16:9", "Figure 2. The main result, as large as it can be."),
         ("text", "Point out what to look at in Figure 2 — the axis, the trend, "
                  "the crossover. Visitors will not find it on their own."),
         ("flex", 3, "free", "Figure 3. Supporting data.")],
        [("heading", "04", "Discussion"),
         ("text", "What the result means, and how it compares with what was "
                  "known. Keep it to the comparison the figure below shows."),
         ("flex", 4, "free", "Figure 4. Comparison with prior work or a control."),
         ("gap",),
         ("conclusion", "05", "Conclusion",
          ["The finding, in one sentence",
           "What it changes for the field",
           "What comes next"]),
         ("gap",),
         ("references", ["[1] A. Author et al., Chem. Mater. 36, 1234 (2024).",
                         "[2] B. Author et al., Adv. Mater. 35, 2301234 (2023).",
                         "[3] C. Author et al., ACS Nano 17, 5678 (2023)."])],
    ]
    report = [P.column(s, x, top, bottom, blocks) for x, blocks in zip(P.col_x, columns)]

    out = Path(out)
    out.parent.mkdir(parents=True, exist_ok=True)
    prs.save(out)
    for i, (end, fh) in enumerate(report, 1):
        print(f"  {i}단 끝 {end:6.1f} mm (본문 하단 {bottom:.1f}) · 가변 그림 높이 {fh:.0f} mm")
    print(f"  {out.name}  {F['name']} {W} × {H} mm · 머리 {HH:.0f} mm")
    return out


if __name__ == "__main__":
    ap = argparse.ArgumentParser()
    ap.add_argument("outdir", nargs="?", default="office")
    ap.add_argument("--format", default="a0-portrait", choices=sorted(FORMATS))
    ap.add_argument("--fonts", help="Archivo-Regular.ttf · Archivo-Bold.ttf 가 있는 폴더")
    a = ap.parse_args()
    build(a.format, Path(a.outdir) / f"hbrlrg-poster-{a.format}.pptx", find_fonts(a.fonts))
